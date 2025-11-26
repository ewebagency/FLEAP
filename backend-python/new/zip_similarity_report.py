from __future__ import annotations

import io
import os
import zipfile
from datetime import datetime
from pathlib import Path
from tempfile import SpooledTemporaryFile
from typing import Dict, List, Optional, Tuple, TypedDict

import fitz  # PyMuPDF
from fastapi import UploadFile
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from new.new_extract_raw import get_raw_text_from_pdf
from new.new_similarity.new_similarity import find_closest_neighbor


class ExpectedNeighbor(TypedDict):
    source: str
    voisin: Optional[str]


class ExtractionPayload(TypedDict):
    text: str
    method: str


class SimilarityRow(TypedDict):
    source: str
    closest_neighbor: Optional[str]
    similarity_score: float
    neighbor_found: bool
    status: str
    expected_neighbor: Optional[str]
    matches_expected: bool
    extraction_method: Optional[str]
    error: Optional[str]


# À mettre à jour en collant le JSON fourni par l'utilisateur.
EXPECTED_NEIGHBORS: List[ExpectedNeighbor] = [
    {
        "source": "facture_ECOBTP_ENVIRONNEMENT_Demathieu_Bard_B_timent_IDF_FACT_-_082025_st_germain_page_1.pdf.pdf",
        "voisin": "facture_Location_de_Bennes_Centre_de_tri_Demathieu_Bard_B_timent_IDF_facture_luxo-1.pdf",
    },
    {
        "source": "facture_RECYCLAGE_DE_L_OUEST__Groupe_Ravate_R2D2__FACT_7182_pages_1_to_2.pdf.pdf",
        "voisin": None,
    },
]


async def build_zip_similarity_report(zip_file: UploadFile) -> Tuple[bytes, List[SimilarityRow]]:
    zip_bytes = await zip_file.read()
    if not zip_bytes:
        raise ValueError("Le fichier ZIP est vide.")

    pdf_payloads, extraction_errors = await _extract_texts_from_zip(zip_bytes)
    if not pdf_payloads:
        raise ValueError("Aucun PDF valide n'a été trouvé dans le ZIP.")

    rows = _compare_all_pdfs(pdf_payloads, extraction_errors)
    excel_bytes = _rows_to_excel(rows)
    return excel_bytes, rows


async def _extract_texts_from_zip(zip_bytes: bytes) -> Tuple[Dict[str, ExtractionPayload], Dict[str, str]]:
    pdf_payloads: Dict[str, ExtractionPayload] = {}
    extraction_errors: Dict[str, str] = {}

    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            if not info.filename.lower().endswith(".pdf"):
                continue

            original_name = Path(info.filename).name or info.filename
            file_key = _ensure_unique_name(original_name, pdf_payloads)
            try:
                pdf_bytes = archive.read(info.filename)
                text, method = await _extract_first_page_text(pdf_bytes, file_key)
                pdf_payloads[file_key] = {"text": text, "method": method}
            except Exception as exc:  # pylint: disable=broad-except
                extraction_errors[file_key] = str(exc)

    return pdf_payloads, extraction_errors


def _ensure_unique_name(candidate: str, existing: Dict[str, ExtractionPayload]) -> str:
    if candidate not in existing:
        return candidate

    stem, ext = os.path.splitext(candidate)
    index = 1
    while True:
        new_name = f"{stem}_{index}{ext}"
        if new_name not in existing:
            return new_name
        index += 1


async def _extract_first_page_text(pdf_bytes: bytes, filename: str) -> Tuple[str, str]:
    first_page_pdf = _extract_first_page_pdf(pdf_bytes)
    temp_file = SpooledTemporaryFile()
    temp_file.write(first_page_pdf)
    temp_file.seek(0)

    upload = UploadFile(file=temp_file, filename=filename, content_type="application/pdf")
    try:
        raw_text, _, parse_or_ocr = await get_raw_text_from_pdf(upload)
        safe_text = raw_text or ""
        return safe_text, parse_or_ocr
    finally:
        await upload.close()
        temp_file.close()


def _extract_first_page_pdf(pdf_bytes: bytes) -> bytes:
    with fitz.open(stream=pdf_bytes, filetype="pdf") as document:
        if document.page_count == 0:
            raise ValueError("PDF sans page.")
        single_page = fitz.open()
        single_page.insert_pdf(document, from_page=0, to_page=0)
        first_page_bytes = single_page.tobytes()
        single_page.close()
    return first_page_bytes


def _compare_all_pdfs(
    pdf_payloads: Dict[str, ExtractionPayload],
    extraction_errors: Dict[str, str],
) -> List[SimilarityRow]:
    expected_lookup = {entry["source"]: entry.get("voisin") for entry in EXPECTED_NEIGHBORS}
    rows: List[SimilarityRow] = []

    for source_name, payload in pdf_payloads.items():
        other_names = [name for name in pdf_payloads if name != source_name]

        if not payload["text"].strip():
            rows.append(
                _build_row(
                    source_name=source_name,
                    closest_neighbor=None,
                    score=0.0,
                    found=False,
                    status="Texte non disponible",
                    expected_neighbor=expected_lookup.get(source_name),
                    extraction_method=payload["method"],
                    error_message=extraction_errors.get(source_name, "Texte vide"),
                )
            )
            continue

        if not other_names:
            rows.append(
                _build_row(
                    source_name=source_name,
                    closest_neighbor=None,
                    score=0.0,
                    found=False,
                    status="Nombre de PDF insuffisant pour comparer",
                    expected_neighbor=expected_lookup.get(source_name),
                    extraction_method=payload["method"],
                    error_message=extraction_errors.get(source_name),
                )
            )
            continue

        neighbor_texts = [pdf_payloads[name]["text"] for name in other_names]
        result = find_closest_neighbor(payload["text"], neighbor_texts, other_names)
        closest_neighbor = result.get("best_neighbor_id")

        rows.append(
            _build_row(
                source_name=source_name,
                closest_neighbor=closest_neighbor,
                score=result["similarity_score"],
                found=result["found"],
                status=result["status"],
                expected_neighbor=expected_lookup.get(source_name),
                extraction_method=payload["method"],
                error_message=extraction_errors.get(source_name),
            )
        )

    return rows


def _build_row(
    source_name: str,
    closest_neighbor: Optional[str],
    score: float,
    found: bool,
    status: str,
    expected_neighbor: Optional[str],
    extraction_method: Optional[str],
    error_message: Optional[str],
) -> SimilarityRow:
    matches_expected = _matches_expectation(
        expected_neighbor=expected_neighbor,
        closest_neighbor=closest_neighbor,
        neighbor_found=found,
    )

    return {
        "source": source_name,
        "closest_neighbor": closest_neighbor,
        "similarity_score": round(score, 4),
        "neighbor_found": found,
        "status": status,
        "expected_neighbor": expected_neighbor,
        "matches_expected": matches_expected,
        "extraction_method": extraction_method,
        "error": error_message,
    }


def _matches_expectation(
    expected_neighbor: Optional[str],
    closest_neighbor: Optional[str],
    neighbor_found: bool,
) -> bool:
    if expected_neighbor is None:
        return not neighbor_found
    return closest_neighbor == expected_neighbor


def _rows_to_excel(rows: List[SimilarityRow]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "similarite_zip"

    headers = [
        "source",
        "closest_neighbor",
        "similarity_score",
        "neighbor_found",
        "status",
        "expected_neighbor",
        "matches_expected",
        "extraction_method",
        "error",
    ]
    worksheet.append(headers)

    for row in rows:
        worksheet.append(
            [
                row["source"],
                row["closest_neighbor"],
                row["similarity_score"],
                row["neighbor_found"],
                row["status"],
                row["expected_neighbor"],
                row["matches_expected"],
                row["extraction_method"],
                row["error"],
            ]
        )

    _auto_fit_columns(worksheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _auto_fit_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        lengths = [len(str(cell.value)) for cell in column_cells if cell.value is not None]
        if not lengths:
            continue
        max_length = max(lengths)
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 80)


def build_report_filename(original_name: Optional[str]) -> str:
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    base = Path(original_name).stem if original_name else "similarity_report"
    sanitized = base.replace(" ", "_") or "similarity_report"
    return f"{sanitized}_{timestamp}.xlsx"

